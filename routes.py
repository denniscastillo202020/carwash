from flask import render_template, request, redirect, url_for, flash, jsonify, send_file
from app import app, db
from models import Customer, Product, ServiceType, Invoice, InvoiceItem, CashRegisterEntry, Appointment, CashClosing
from forms import CustomerForm, ProductForm, ServiceTypeForm, InvoiceForm, CashRegisterForm, AppointmentForm
from utils import (generate_invoice_number, format_lempiras, calculate_cash_total, 
                   send_whatsapp_message, get_available_time_slots, perform_cash_closing, update_customer_crm)
from datetime import datetime, date, timedelta
from sqlalchemy import func, or_
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io

@app.route('/')
def index():
    """Dashboard with key metrics"""
    today = datetime.now().date()
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())
    
    # Today's metrics
    today_sales = db.session.query(func.sum(Invoice.total_amount)).filter(
        Invoice.created_at.between(today_start, today_end)
    ).scalar() or 0.0
    
    today_invoices = Invoice.query.filter(
        Invoice.created_at.between(today_start, today_end)
    ).count()
    
    today_appointments = Appointment.query.filter(
        func.date(Appointment.appointment_date) == today,
        Appointment.status == 'scheduled'
    ).count()
    
    # Low stock products
    low_stock_products = Product.query.filter(Product.stock <= 5).all()
    
    # Recent customers
    recent_customers = Customer.query.order_by(Customer.created_at.desc()).limit(5).all()
    
    return render_template('enhanced_index.html',
                         daily_sales=today_sales,
                         daily_invoices=today_invoices,
                         daily_appointments=today_appointments,
                         monthly_invoices=today_invoices * 30,
                         monthly_sales=today_sales * 30,
                         total_customers=Customer.query.count(),
                         new_customers_month=0,
                         pending_appointments=today_appointments,
                         recent_invoices=[],
                         today_appointments=[],
                         low_stock_products=low_stock_products)

@app.route('/inventory')
def inventory():
    """Inventory management"""
    products = Product.query.all()
    services = ServiceType.query.all()
    return render_template('inventory.html', products=products, services=services)

@app.route('/add_product', methods=['GET', 'POST'])
def add_product():
    """Add new product"""
    form = ProductForm()
    if form.validate_on_submit():
        # Check if barcode already exists
        if form.barcode.data and Product.query.filter_by(barcode=form.barcode.data).first():
            flash('El código de barras ya existe', 'error')
            return render_template('inventory.html', form=form)
        
        product = Product(
            name=form.name.data,
            barcode=form.barcode.data,
            price=form.price.data,
            stock=form.stock.data,
            category=form.category.data,
            description=form.description.data
        )
        
        db.session.add(product)
        db.session.commit()
        flash('Producto agregado exitosamente', 'success')
        return redirect(url_for('inventory'))
    
    # If form validation fails, re-render with product list
    products = Product.query.all()
    services = ServiceType.query.all()
    return render_template('inventory.html', products=products, services=services, form=form)

@app.route('/add_service', methods=['GET', 'POST'])
def add_service():
    """Add new service type"""
    form = ServiceTypeForm()
    if form.validate_on_submit():
        service = ServiceType(
            name=form.name.data,
            price=form.price.data,
            category=form.category.data,
            duration_minutes=form.duration_minutes.data,
            description=form.description.data
        )
        
        db.session.add(service)
        db.session.commit()
        flash('Servicio agregado exitosamente', 'success')
        return redirect(url_for('inventory'))
    
    # If form validation fails, re-render with service list
    products = Product.query.all()
    services = ServiceType.query.all()
    return render_template('inventory.html', products=products, services=services, form=form)

@app.route('/edit_product', methods=['POST'])
def edit_product():
    """Edit existing product"""
    try:
        product_id = request.form.get('product_id')
        product = Product.query.get_or_404(product_id)
        
        # Update product fields
        product.name = request.form.get('name')
        product.barcode = request.form.get('barcode') or None
        product.price = float(request.form.get('price'))
        product.stock = int(request.form.get('stock'))
        product.category = request.form.get('category')
        
        db.session.commit()
        flash('Producto actualizado exitosamente', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error al actualizar producto: {str(e)}', 'error')
        
    return redirect(url_for('inventory'))

@app.route('/customers')
def customers():
    """Customer management (CRM)"""
    search = request.args.get('search', '')
    if search:
        customers_list = Customer.query.filter(
            or_(Customer.name.contains(search),
                Customer.phone.contains(search),
                Customer.email.contains(search))
        ).order_by(Customer.created_at.desc()).all()
    else:
        customers_list = Customer.query.order_by(Customer.created_at.desc()).all()
    
    return render_template('customers.html', customers=customers_list, search=search)

@app.route('/add_customer', methods=['GET', 'POST'])
def add_customer():
    """Add new customer"""
    form = CustomerForm()
    if form.validate_on_submit():
        # Check if customer already exists
        existing_customer = Customer.query.filter_by(phone=form.phone.data).first()
        if existing_customer:
            flash('Ya existe un cliente con este número de teléfono', 'error')
            return render_template('customers.html', form=form)
        
        customer = Customer(
            name=form.name.data,
            phone=form.phone.data,
            email=form.email.data,
            address=form.address.data
        )
        
        db.session.add(customer)
        db.session.commit()
        flash('Cliente agregado exitosamente', 'success')
        return redirect(url_for('customers'))
    
    # If form validation fails, re-render with customers list
    customers_list = Customer.query.order_by(Customer.created_at.desc()).all()
    return render_template('customers.html', customers=customers_list, form=form)

@app.route('/invoice')
def invoice():
    """Invoice creation"""
    products = Product.query.filter(Product.stock > 0).all()
    services = ServiceType.query.all()
    return render_template('simple_invoice.html', products=products, services=services, now=datetime.now())

@app.route('/create_invoice', methods=['POST'])
def create_invoice():
    """Create new invoice"""
    try:
        # Get customer data
        customer_name = request.form.get('customer_name', '').strip()
        customer_phone = request.form.get('customer_phone', '').strip()
        customer_email = request.form.get('customer_email', '').strip()
        customer_address = request.form.get('customer_address', '').strip()
        payment_method = request.form.get('payment_method', 'efectivo')
        
        if not customer_name or not customer_phone:
            flash('Nombre y teléfono del cliente son obligatorios', 'error')
            return redirect(url_for('invoice'))
        
        # Find or create customer
        customer = Customer.query.filter_by(phone=customer_phone).first()
        if not customer:
            customer = Customer(
                name=customer_name,
                phone=customer_phone,
                email=customer_email,
                address=customer_address
            )
            db.session.add(customer)
            db.session.flush()  # Get customer ID
        else:
            # Update customer info if needed
            customer.name = customer_name
            customer.email = customer_email or customer.email
            customer.address = customer_address or customer.address
        
        # Create invoice
        invoice = Invoice(
            invoice_number=generate_invoice_number(),
            customer_id=customer.id,
            total_amount=0.0,
            payment_method=payment_method
        )
        db.session.add(invoice)
        db.session.flush()  # Get invoice ID
        
        # Process items
        total_amount = 0.0
        items_processed = False
        
        # Products
        for key in request.form:
            if key.startswith('product_quantity_'):
                product_id = key.split('_')[-1]
                quantity = int(request.form.get(key, 0))
                
                if quantity > 0:
                    product = Product.query.get(product_id)
                    if product and product.stock >= quantity:
                        item_total = product.price * quantity
                        
                        invoice_item = InvoiceItem(
                            invoice_id=invoice.id,
                            product_id=product.id,
                            quantity=quantity,
                            unit_price=product.price,
                            total_price=item_total,
                            description=product.name
                        )
                        db.session.add(invoice_item)
                        
                        # Update stock
                        product.stock -= quantity
                        total_amount += item_total
                        items_processed = True
        
        # Services
        for key in request.form:
            if key.startswith('service_quantity_'):
                service_id = key.split('_')[-1]
                quantity = int(request.form.get(key, 0))
                
                if quantity > 0:
                    service = ServiceType.query.get(service_id)
                    if service:
                        item_total = service.price * quantity
                        
                        invoice_item = InvoiceItem(
                            invoice_id=invoice.id,
                            service_type_id=service.id,
                            quantity=quantity,
                            unit_price=service.price,
                            total_price=item_total,
                            description=service.name
                        )
                        db.session.add(invoice_item)
                        total_amount += item_total
                        items_processed = True
        
        if not items_processed:
            flash('Debe agregar al menos un producto o servicio', 'error')
            db.session.rollback()
            return redirect(url_for('invoice'))
        
        # Update invoice total
        invoice.total_amount = total_amount
        
        # Process cash payment if efectivo
        if payment_method == 'efectivo':
            # Get denomination data from form
            cash_data = {}
            for field in ['bills_500', 'bills_200', 'bills_100', 'bills_50', 'bills_20', 
                         'bills_10', 'bills_5', 'bills_2', 'bills_1', 
                         'coins_50c', 'coins_20c', 'coins_10c', 'coins_5c']:
                cash_data[field] = int(request.form.get(field, 0))
            
            # Calculate total cash received
            total_cash = calculate_cash_total(cash_data)
            
            # Create cash register entry
            cash_entry = CashRegisterEntry(
                invoice_id=invoice.id,
                total_amount=total_cash,
                **cash_data
            )
            db.session.add(cash_entry)
        
        # Update customer CRM
        update_customer_crm(customer.id)
        
        db.session.commit()
        flash(f'Factura {invoice.invoice_number} creada exitosamente', 'success')
        return redirect(url_for('print_invoice', invoice_id=invoice.id))
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error al crear factura: {str(e)}', 'error')
        return redirect(url_for('invoice'))

@app.route('/cash_register/<int:invoice_id>')
def cash_register(invoice_id):
    """Cash register for invoice payment"""
    invoice = Invoice.query.get_or_404(invoice_id)
    return render_template('simple_cash_register.html', invoice=invoice)

@app.route('/process_cash_payment/<int:invoice_id>', methods=['POST'])
def process_cash_payment(invoice_id):
    """Process cash payment with denominations"""
    invoice = Invoice.query.get_or_404(invoice_id)
    
    try:
        # Get cash data directly from form
        cash_data = {

            'bills_500': int(request.form.get('bills_500', 0)),
            'bills_200': int(request.form.get('bills_200', 0)),
            'bills_100': int(request.form.get('bills_100', 0)),
            'bills_50': int(request.form.get('bills_50', 0)),
            'bills_20': int(request.form.get('bills_20', 0)),
            'bills_10': int(request.form.get('bills_10', 0)),
            'bills_5': int(request.form.get('bills_5', 0)),
            'bills_2': int(request.form.get('bills_2', 0)),
            'bills_1': int(request.form.get('bills_1', 0)),
            'coins_50c': int(request.form.get('coins_50c', 0)),
            'coins_20c': int(request.form.get('coins_20c', 0)),
            'coins_10c': int(request.form.get('coins_10c', 0)),
            'coins_5c': int(request.form.get('coins_5c', 0)),
        }
        
        total_received = calculate_cash_total(cash_data)
        
        if total_received < invoice.total_amount:
            flash(f'Cantidad insuficiente. Se requieren {format_lempiras(invoice.total_amount)}', 'error')
            return render_template('simple_cash_register.html', invoice=invoice)
        
        # Record cash register entry
        cash_entry = CashRegisterEntry(
            invoice_id=invoice.id,
            **cash_data,
            total_amount=total_received
        )
        
        db.session.add(cash_entry)
        
        # Update customer CRM
        update_customer_crm(invoice.customer_id)
        
        db.session.commit()
        
        change = total_received - invoice.total_amount
        if change > 0:
            flash(f'Pago procesado. Cambio: {format_lempiras(change)}', 'success')
        else:
            flash('Pago procesado exitosamente', 'success')
        
        return redirect(url_for('print_invoice', invoice_id=invoice.id))
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error al procesar pago: {str(e)}', 'error')
        return render_template('simple_cash_register.html', invoice=invoice)
    


@app.route('/cash_closing')
def cash_closing():
    """Daily cash closing"""
    today = datetime.now().date()
    existing_closing = CashClosing.query.filter_by(closing_date=today).first()
    
    # Calculate today's sales by payment method
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())
    
    # Cash sales
    cash_sales = db.session.query(func.sum(Invoice.total_amount)).filter(
        Invoice.created_at.between(today_start, today_end),
        Invoice.payment_method == 'efectivo'
    ).scalar() or 0.0
    
    cash_invoices = Invoice.query.filter(
        Invoice.created_at.between(today_start, today_end),
        Invoice.payment_method == 'efectivo'
    ).count()
    
    # Card sales
    card_sales = db.session.query(func.sum(Invoice.total_amount)).filter(
        Invoice.created_at.between(today_start, today_end),
        Invoice.payment_method == 'tarjeta'
    ).scalar() or 0.0
    
    card_invoices = Invoice.query.filter(
        Invoice.created_at.between(today_start, today_end),
        Invoice.payment_method == 'tarjeta'
    ).count()
    
    # Total sales
    total_sales = cash_sales + card_sales
    total_invoices = cash_invoices + card_invoices
    
    # Get current cash total from all cash entries today
    cash_entries = CashRegisterEntry.query.join(Invoice).filter(
        Invoice.created_at.between(today_start, today_end)
    ).all()
    
    # Sum up all cash denominations
    current_cash = {
        'bills_500': sum(entry.bills_500 for entry in cash_entries),
        'bills_200': sum(entry.bills_200 for entry in cash_entries),
        'bills_100': sum(entry.bills_100 for entry in cash_entries),
        'bills_50': sum(entry.bills_50 for entry in cash_entries),
        'bills_20': sum(entry.bills_20 for entry in cash_entries),
        'bills_10': sum(entry.bills_10 for entry in cash_entries),
        'bills_5': sum(entry.bills_5 for entry in cash_entries),
        'bills_2': sum(entry.bills_2 for entry in cash_entries),
        'bills_1': sum(entry.bills_1 for entry in cash_entries),
    }
    
    # Create a simple object to access in template
    class CashSummary:
        def __init__(self, data):
            for key, value in data.items():
                setattr(self, key, value)
    
    current_cash_obj = CashSummary(current_cash)
    
    return render_template('simple_cash_closing.html',
                         today=today,
                         existing_closing=existing_closing,
                         cash_sales=cash_sales,
                         cash_invoices=cash_invoices,
                         card_sales=card_sales,
                         card_invoices=card_invoices,
                         total_sales=total_sales,
                         total_invoices=total_invoices,
                         current_cash=current_cash_obj)

@app.route('/perform_cash_closing', methods=['POST'])
def perform_cash_closing_route():
    """Perform daily cash closing"""
    try:
        today = datetime.now().date()
        
        # Check if closing already exists
        existing_closing = CashClosing.query.filter_by(closing_date=today).first()
        if existing_closing:
            flash('Ya se realizó el cierre de caja para hoy', 'warning')
            return redirect(url_for('cash_closing'))
        
        # Get cash data from form
        cash_data = {
            'total_bills_500': int(request.form.get('bills_500', 0)),
            'total_bills_200': int(request.form.get('bills_200', 0)),
            'total_bills_100': int(request.form.get('bills_100', 0)),
            'total_bills_50': int(request.form.get('bills_50', 0)),
            'total_bills_20': int(request.form.get('bills_20', 0)),
            'total_bills_10': int(request.form.get('bills_10', 0)),
            'total_bills_5': int(request.form.get('bills_5', 0)),
            'total_bills_2': int(request.form.get('bills_2', 0)),
            'total_bills_1': int(request.form.get('bills_1', 0)),
        }
        
        expected_amount = float(request.form.get('expected_amount', 0))
        notes = request.form.get('notes', '')
        
        # Calculate actual amount
        actual_amount = (
            cash_data['total_bills_500'] * 500 +
            cash_data['total_bills_200'] * 200 +
            cash_data['total_bills_100'] * 100 +
            cash_data['total_bills_50'] * 50 +
            cash_data['total_bills_20'] * 20 +
            cash_data['total_bills_10'] * 10 +
            cash_data['total_bills_5'] * 5 +
            cash_data['total_bills_2'] * 2 +
            cash_data['total_bills_1'] * 1
        )
        
        difference = actual_amount - expected_amount
        
        # Create cash closing record
        cash_closing = CashClosing(
            closing_date=today,
            sales_amount=expected_amount,
            expected_amount=expected_amount,
            actual_amount=actual_amount,
            difference=difference,
            notes=notes,
            **cash_data
        )
        
        db.session.add(cash_closing)
        db.session.commit()
        
        if difference == 0:
            flash('Cierre de caja perfecto. Sin diferencias.', 'success')
        elif difference > 0:
            flash(f'Cierre realizado. Sobrante: {format_lempiras(difference)}', 'warning')
        else:
            flash(f'Cierre realizado. Faltante: {format_lempiras(abs(difference))}', 'danger')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error al realizar cierre: {str(e)}', 'error')
    
    return redirect(url_for('cash_closing'))

@app.route('/analytics')
def analytics():
    """Analytics dashboard"""
    return render_template('analytics.html')

@app.route('/api/analytics', methods=['POST'])
def api_analytics():
    """API endpoint for analytics data"""
    try:
        period = request.form.get('period', 'week')
        date_from = request.form.get('date_from')
        date_to = request.form.get('date_to')
        
        # Calculate date range based on period
        today = datetime.now().date()
        if period == 'today':
            start_date = today
            end_date = today
        elif period == 'week':
            start_date = today - timedelta(days=today.weekday())
            end_date = start_date + timedelta(days=6)
        elif period == 'month':
            start_date = today.replace(day=1)
            next_month = start_date.replace(month=start_date.month + 1) if start_date.month < 12 else start_date.replace(year=start_date.year + 1, month=1)
            end_date = next_month - timedelta(days=1)
        elif period == 'quarter':
            quarter = (today.month - 1) // 3 + 1
            start_date = datetime(today.year, (quarter - 1) * 3 + 1, 1).date()
            if quarter < 4:
                end_date = datetime(today.year, quarter * 3, 1).date()
                end_date = end_date.replace(month=end_date.month + 1) - timedelta(days=1)
            else:
                end_date = datetime(today.year, 12, 31).date()
        elif period == 'year':
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)
        elif period == 'custom' and date_from and date_to:
            start_date = datetime.strptime(date_from, '%Y-%m-%d').date()
            end_date = datetime.strptime(date_to, '%Y-%m-%d').date()
        else:
            start_date = today - timedelta(days=7)
            end_date = today
        
        # Convert to datetime for queries
        start_datetime = datetime.combine(start_date, datetime.min.time())
        end_datetime = datetime.combine(end_date, datetime.max.time())
        
        # Get invoices for the period
        invoices = Invoice.query.filter(
            Invoice.created_at.between(start_datetime, end_datetime)
        ).all()
        
        # Calculate KPIs
        total_sales = sum(inv.total_amount for inv in invoices)
        total_invoices = len(invoices)
        unique_customers = len(set(inv.customer_id for inv in invoices))
        avg_ticket = total_sales / total_invoices if total_invoices > 0 else 0
        
        # Calculate previous period for growth
        period_days = (end_date - start_date).days + 1
        prev_start = start_date - timedelta(days=period_days)
        prev_end = start_date - timedelta(days=1)
        prev_start_datetime = datetime.combine(prev_start, datetime.min.time())
        prev_end_datetime = datetime.combine(prev_end, datetime.max.time())
        
        prev_invoices = Invoice.query.filter(
            Invoice.created_at.between(prev_start_datetime, prev_end_datetime)
        ).all()
        
        prev_sales = sum(inv.total_amount for inv in prev_invoices)
        prev_invoices_count = len(prev_invoices)
        prev_customers = len(set(inv.customer_id for inv in prev_invoices))
        prev_avg_ticket = prev_sales / prev_invoices_count if prev_invoices_count > 0 else 0
        
        # Calculate growth percentages
        sales_growth = ((total_sales - prev_sales) / prev_sales * 100) if prev_sales > 0 else 0
        invoices_growth = ((total_invoices - prev_invoices_count) / prev_invoices_count * 100) if prev_invoices_count > 0 else 0
        customers_growth = ((unique_customers - prev_customers) / prev_customers * 100) if prev_customers > 0 else 0
        ticket_growth = ((avg_ticket - prev_avg_ticket) / prev_avg_ticket * 100) if prev_avg_ticket > 0 else 0
        
        # Sales trend data
        sales_trend_labels = []
        sales_trend_data = []
        current_date = start_date
        while current_date <= end_date:
            day_start = datetime.combine(current_date, datetime.min.time())
            day_end = datetime.combine(current_date, datetime.max.time())
            day_sales = sum(inv.total_amount for inv in invoices 
                          if day_start <= inv.created_at <= day_end)
            sales_trend_labels.append(current_date.strftime('%d/%m'))
            sales_trend_data.append(float(day_sales))
            current_date += timedelta(days=1)
        
        # Payment methods
        cash_sales = sum(inv.total_amount for inv in invoices if inv.payment_method == 'efectivo')
        card_sales = sum(inv.total_amount for inv in invoices if inv.payment_method == 'tarjeta')
        transfer_sales = sum(inv.total_amount for inv in invoices if inv.payment_method == 'transferencia')
        
        # Top services and products
        from collections import defaultdict
        service_counts = defaultdict(int)
        product_counts = defaultdict(int)
        
        for invoice in invoices:
            for item in invoice.items:
                if item.service_type:
                    service_counts[item.service_type.name] += item.quantity
                if item.product:
                    product_counts[item.product.name] += item.quantity
        
        top_services = sorted(service_counts.items(), key=lambda x: x[1], reverse=True)[:5]
        top_products = sorted(product_counts.items(), key=lambda x: x[1], reverse=True)[:5]
        
        # Hourly analysis
        hourly_sales = [0] * 13  # 6AM to 6PM
        for invoice in invoices:
            hour = invoice.created_at.hour
            if 6 <= hour <= 18:
                hourly_sales[hour - 6] += invoice.total_amount
        
        # Weekly analysis
        weekly_sales = [0] * 7  # Monday to Sunday
        for invoice in invoices:
            weekday = invoice.created_at.weekday()
            weekly_sales[weekday] += invoice.total_amount
        
        # Top customers
        from collections import defaultdict
        customer_stats = defaultdict(lambda: {'total': 0, 'count': 0, 'last_visit': None})
        for invoice in invoices:
            customer_stats[invoice.customer_id]['total'] += invoice.total_amount
            customer_stats[invoice.customer_id]['count'] += 1
            if not customer_stats[invoice.customer_id]['last_visit'] or invoice.created_at > customer_stats[invoice.customer_id]['last_visit']:
                customer_stats[invoice.customer_id]['last_visit'] = invoice.created_at
        
        top_customers = []
        for customer_id, stats in customer_stats.items():
            customer = Customer.query.get(customer_id)
            if customer:
                top_customers.append({
                    'name': customer.name,
                    'phone': customer.phone,
                    'total_spent': stats['total'],
                    'invoice_count': stats['count'],
                    'avg_ticket': stats['total'] / stats['count'],
                    'last_visit': stats['last_visit'].strftime('%d/%m/%Y')
                })
        
        top_customers.sort(key=lambda x: x['total_spent'], reverse=True)
        top_customers = top_customers[:10]
        
        return jsonify({
            'kpis': {
                'total_sales': round(total_sales, 2),
                'total_invoices': total_invoices,
                'unique_customers': unique_customers,
                'avg_ticket': round(avg_ticket, 2),
                'sales_growth': round(sales_growth, 1),
                'invoices_growth': round(invoices_growth, 1),
                'customers_growth': round(customers_growth, 1),
                'ticket_growth': round(ticket_growth, 1)
            },
            'charts': {
                'sales_trend': {
                    'labels': sales_trend_labels,
                    'data': sales_trend_data
                },
                'payment_methods': {
                    'data': [float(cash_sales), float(card_sales), float(transfer_sales)]
                },
                'top_services': {
                    'labels': [item[0] for item in top_services],
                    'data': [item[1] for item in top_services]
                },
                'top_products': {
                    'labels': [item[0] for item in top_products],
                    'data': [item[1] for item in top_products]
                },
                'hourly': {
                    'data': [float(x) for x in hourly_sales]
                },
                'weekly': {
                    'data': [float(x) for x in weekly_sales]
                }
            },
            'top_customers': top_customers
        })
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/export_cash_closings')
def export_cash_closings():
    """Export all cash closings to Excel"""
    try:
        import xlsxwriter
        from io import BytesIO
        
        # Create a workbook and add a worksheet
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output)
        worksheet = workbook.add_worksheet('Arqueos de Caja')
        
        # Define formats
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#D7E4BC',
            'border': 1
        })
        
        money_format = workbook.add_format({
            'num_format': 'L#,##0.00',
            'border': 1
        })
        
        date_format = workbook.add_format({
            'num_format': 'dd/mm/yyyy',
            'border': 1
        })
        
        cell_format = workbook.add_format({'border': 1})
        
        # Headers
        headers = [
            'Fecha', 'Ventas Esperadas', 'Efectivo Contado', 'Diferencia',
            'L500', 'L200', 'L100', 'L50', 'L20', 'L10', 'L5', 'L2', 'L1',
            'Notas', 'Fecha Creación'
        ]
        
        for col, header in enumerate(headers):
            worksheet.write(0, col, header, header_format)
        
        # Get all cash closings
        closings = CashClosing.query.order_by(CashClosing.closing_date.desc()).all()
        
        # Write data
        for row, closing in enumerate(closings, 1):
            worksheet.write(row, 0, closing.closing_date, date_format)
            worksheet.write(row, 1, closing.expected_amount, money_format)
            worksheet.write(row, 2, closing.actual_amount, money_format)
            worksheet.write(row, 3, closing.difference, money_format)
            worksheet.write(row, 4, closing.total_bills_500 or 0, cell_format)
            worksheet.write(row, 5, closing.total_bills_200 or 0, cell_format)
            worksheet.write(row, 6, closing.total_bills_100 or 0, cell_format)
            worksheet.write(row, 7, closing.total_bills_50 or 0, cell_format)
            worksheet.write(row, 8, closing.total_bills_20 or 0, cell_format)
            worksheet.write(row, 9, closing.total_bills_10 or 0, cell_format)
            worksheet.write(row, 10, closing.total_bills_5 or 0, cell_format)
            worksheet.write(row, 11, closing.total_bills_2 or 0, cell_format)
            worksheet.write(row, 12, closing.total_bills_1 or 0, cell_format)
            worksheet.write(row, 13, closing.notes or '', cell_format)
            worksheet.write(row, 14, closing.created_at, date_format)
        
        # Adjust column widths
        worksheet.set_column('A:A', 12)
        worksheet.set_column('B:D', 15)
        worksheet.set_column('E:M', 8)
        worksheet.set_column('N:N', 20)
        worksheet.set_column('O:O', 15)
        
        workbook.close()
        output.seek(0)
        
        filename = f'Arqueos_Caja_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        flash(f'Error al exportar arqueos: {str(e)}', 'error')
        return redirect(url_for('cash_closing'))

@app.route('/export_analytics')
def export_analytics():
    """Export analytics report to Excel"""
    try:
        import xlsxwriter
        from io import BytesIO
        
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output)
        
        # Create sheets
        summary_sheet = workbook.add_worksheet('Resumen')
        daily_sheet = workbook.add_worksheet('Ventas Diarias')
        customers_sheet = workbook.add_worksheet('Top Clientes')
        
        # Formats
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#4472C4',
            'font_color': 'white',
            'border': 1
        })
        
        money_format = workbook.add_format({
            'num_format': 'L#,##0.00',
            'border': 1
        })
        
        # Get current month data
        today = datetime.now().date()
        start_date = today.replace(day=1)
        end_date = today
        start_datetime = datetime.combine(start_date, datetime.min.time())
        end_datetime = datetime.combine(end_date, datetime.max.time())
        
        invoices = Invoice.query.filter(
            Invoice.created_at.between(start_datetime, end_datetime)
        ).all()
        
        # Summary sheet
        summary_sheet.write('A1', 'REPORTE DE ANÁLISIS - CAR WASH PEÑA BLANCA', header_format)
        summary_sheet.merge_range('A1:D1', 'REPORTE DE ANÁLISIS - CAR WASH PEÑA BLANCA', header_format)
        
        summary_data = [
            ['Período', f'{start_date.strftime("%d/%m/%Y")} - {end_date.strftime("%d/%m/%Y")}'],
            ['Total Ventas', sum(inv.total_amount for inv in invoices)],
            ['Total Facturas', len(invoices)],
            ['Clientes Únicos', len(set(inv.customer_id for inv in invoices))],
            ['Ticket Promedio', sum(inv.total_amount for inv in invoices) / len(invoices) if invoices else 0],
            ['Ventas Efectivo', sum(inv.total_amount for inv in invoices if inv.payment_method == 'efectivo')],
            ['Ventas Tarjeta', sum(inv.total_amount for inv in invoices if inv.payment_method == 'tarjeta')]
        ]
        
        for row, (label, value) in enumerate(summary_data, 3):
            summary_sheet.write(row, 0, label, header_format)
            if isinstance(value, (int, float)) and 'Ventas' in label or 'Promedio' in label:
                summary_sheet.write(row, 1, value, money_format)
            else:
                summary_sheet.write(row, 1, value)
        
        workbook.close()
        output.seek(0)
        
        filename = f'Reporte_Analytics_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        flash(f'Error al exportar reporte: {str(e)}', 'error')
        return redirect(url_for('analytics'))

@app.route('/appointments')
def appointments():
    """Simplified appointment management"""
    # Get services for dropdown
    services = ServiceType.query.all()
    
    # Get today's appointments
    today = datetime.now().date()
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())
    
    todays_appointments = Appointment.query.filter(
        Appointment.appointment_date.between(today_start, today_end)
    ).order_by(Appointment.appointment_date).all()
    
    return render_template('simple_appointments.html', 
                         services=services,
                         todays_appointments=todays_appointments,
                         today=today)

@app.route('/add_appointment', methods=['POST'])
def add_appointment():
    """Add new appointment - simplified version"""
    try:
        # Get form data
        customer_phone = request.form.get('customer_phone', '').strip()
        customer_name = request.form.get('customer_name', '').strip()
        appointment_date = request.form.get('appointment_date')
        appointment_time = request.form.get('appointment_time')
        service_type_id = request.form.get('service_type_id')
        notes = request.form.get('notes', '').strip()
        
        # Validate required fields
        if not all([customer_phone, customer_name, appointment_date, appointment_time, service_type_id]):
            flash('Todos los campos son obligatorios', 'error')
            return redirect(url_for('appointments'))
        
        # Combine date and time
        appointment_datetime = datetime.strptime(f"{appointment_date} {appointment_time}", '%Y-%m-%d %H:%M')
        
        # Find or create customer
        customer = Customer.query.filter_by(phone=customer_phone).first()
        if not customer:
            customer = Customer(
                name=customer_name,
                phone=customer_phone
            )
            db.session.add(customer)
            db.session.flush()
        else:
            # Update customer name if different
            customer.name = customer_name
        
        # Create appointment
        appointment = Appointment(
            customer_id=customer.id,
            service_type_id=int(service_type_id),
            appointment_date=appointment_datetime,
            notes=notes,
            status='scheduled'
        )
        
        db.session.add(appointment)
        db.session.commit()
        
        # Send WhatsApp confirmation
        service = ServiceType.query.get(int(service_type_id))
        message = f"""✅ *CITA CONFIRMADA*

Hola {customer.name},

Tu cita ha sido agendada:
📅 *Fecha:* {appointment_datetime.strftime('%d/%m/%Y')}
🕐 *Hora:* {appointment_datetime.strftime('%H:%M')}
🚗 *Servicio:* {service.name}
💰 *Precio:* {format_lempiras(service.price)}

📍 *Car Wash Peña Blanca*
📞 Tel: 97164446

¡Te esperamos! 🚗✨"""
        
        send_whatsapp_message(customer.phone, message)
        
        flash(f'Cita agendada para {customer.name} - Confirmación enviada por WhatsApp', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error al agendar cita: {str(e)}', 'error')
    
    return redirect(url_for('appointments'))

@app.route('/complete_appointment/<int:appointment_id>', methods=['POST'])
def complete_appointment(appointment_id):
    """Mark appointment as completed"""
    try:
        appointment = Appointment.query.get_or_404(appointment_id)
        appointment.status = 'completed'
        db.session.commit()
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)})

@app.route('/send_ready_notification/<int:appointment_id>', methods=['POST'])
def send_ready_notification(appointment_id):
    """Send WhatsApp notification that car is ready"""
    try:
        appointment = Appointment.query.get_or_404(appointment_id)
        customer = appointment.customer
        
        message = f"""🚗✨ *¡TU VEHÍCULO ESTÁ LISTO!*

Hola {customer.name},

Tu {appointment.service_type.name} ha sido completado.

📍 *Car Wash Peña Blanca*
📞 Tel: 97164446

¡Ven a recoger tu vehículo reluciente! 🌟"""
        
        send_whatsapp_message(customer.phone, message)
        return jsonify({'success': True})
        
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

@app.route('/whatsapp_webhook', methods=['GET', 'POST'])
def whatsapp_webhook():
    """WhatsApp Bot webhook for automated responses"""
    if request.method == 'GET':
        # Webhook verification
        verify_token = "car_wash_bot_token"
        mode = request.args.get('hub.mode')
        token = request.args.get('hub.verify_token')
        challenge = request.args.get('hub.challenge')
        
        if mode and token:
            if mode == 'subscribe' and token == verify_token:
                return challenge
        return 'Forbidden', 403
    
    elif request.method == 'POST':
        # Process incoming messages
        try:
            body = request.get_json()
            
            if body.get('object') == 'whatsapp_business_account':
                for entry in body.get('entry', []):
                    for change in entry.get('changes', []):
                        if change.get('field') == 'messages':
                            messages = change.get('value', {}).get('messages', [])
                            
                            for message in messages:
                                phone = message.get('from')
                                text = message.get('text', {}).get('body', '').lower()
                                
                                # Process bot commands
                                if any(word in text for word in ['hola', 'cita', 'agendar', 'horarios']):
                                    handle_whatsapp_message(phone, text)
            
            return 'OK', 200
            
        except Exception as e:
            return f'Error: {str(e)}', 500

def handle_whatsapp_message(phone, text):
    """Handle incoming WhatsApp messages"""
    try:
        if 'horarios' in text or 'disponibles' in text:
            # Send available time slots
            response = get_available_slots_message()
            send_whatsapp_message(phone, response)
            
        elif 'agendar' in text:
            # Extract date and time from message
            # Simple parsing - could be enhanced
            response = """Para agendar una cita, envía un mensaje como:
*"Agendar mañana 10:00"* o
*"Agendar 15/01 14:00"*

O llámanos al 97164446 📞"""
            send_whatsapp_message(phone, response)
            
        else:
            # Default greeting and options
            response = f"""¡Hola! 👋 Bienvenido a *Car Wash Peña Blanca*

🚗 *Servicios disponibles:*
• Lavado Básico - L80
• Lavado Premium - L120  
• Encerado - L150

📅 *Para agendar:*
Escribe "horarios disponibles" o llámanos al 97164446

📍 Peña Blanca, Cortés
🕐 Lunes a Sábado: 7:00 AM - 6:00 PM"""
            
            send_whatsapp_message(phone, response)
            
    except Exception as e:
        print(f"Error handling WhatsApp message: {e}")

def get_available_slots_message():
    """Generate message with available time slots"""
    try:
        today = datetime.now().date()
        tomorrow = today + timedelta(days=1)
        
        # Get busy slots for today and tomorrow
        today_appointments = Appointment.query.filter(
            func.date(Appointment.appointment_date) == today,
            Appointment.status == 'scheduled'
        ).all()
        
        tomorrow_appointments = Appointment.query.filter(
            func.date(Appointment.appointment_date) == tomorrow,
            Appointment.status == 'scheduled'
        ).all()
        
        # Define working hours
        working_hours = ['07:00', '08:00', '09:00', '10:00', '11:00', '12:00', 
                        '13:00', '14:00', '15:00', '16:00', '17:00']
        
        # Get available slots for today
        today_busy = [apt.appointment_date.strftime('%H:%M') for apt in today_appointments]
        today_available = [hour for hour in working_hours if hour not in today_busy]
        
        # Get available slots for tomorrow  
        tomorrow_busy = [apt.appointment_date.strftime('%H:%M') for apt in tomorrow_appointments]
        tomorrow_available = [hour for hour in working_hours if hour not in tomorrow_busy]
        
        message = f"""📅 *HORARIOS DISPONIBLES*

*HOY ({today.strftime('%d/%m')}):*
{', '.join(today_available) if today_available else 'No hay horarios disponibles'}

*MAÑANA ({tomorrow.strftime('%d/%m')}):*
{', '.join(tomorrow_available) if tomorrow_available else 'No hay horarios disponibles'}

Para agendar llama al 97164446 📞"""
        
        return message
        
    except Exception as e:
        return "Error al consultar horarios. Llama al 97164446 📞"

@app.route('/api/available_slots')
def api_available_slots():
    """API endpoint for available time slots"""
    date_str = request.args.get('date')
    service_id = request.args.get('service_id')
    
    if not date_str or not service_id:
        return jsonify({'error': 'Date and service_id required'}), 400
    
    try:
        date_obj = datetime.strptime(date_str, '%Y-%m-%d').date()
        service = ServiceType.query.get(service_id)
        
        if not service:
            return jsonify({'error': 'Service not found'}), 404
        
        slots = get_available_time_slots(date_obj, service.duration_minutes)
        return jsonify({'slots': slots})
        
    except ValueError:
        return jsonify({'error': 'Invalid date format'}), 400

@app.route('/search_customer')
def search_customer():
    """Search customer by phone for autocomplete"""
    phone = request.args.get('phone', '')
    if len(phone) >= 4:
        customer = Customer.query.filter_by(phone=phone).first()
        if customer:
            return jsonify({
                'found': True,
                'name': customer.name,
                'email': customer.email or '',
                'address': customer.address or '',
                'total_visits': customer.total_visits or 0
            })
    
    return jsonify({'found': False})

@app.route('/search_product')
def search_product():
    """Search product by barcode"""
    barcode = request.args.get('barcode', '')
    if barcode:
        product = Product.query.filter_by(barcode=barcode).first()
        if product:
            return jsonify({
                'found': True,
                'id': product.id,
                'name': product.name,
                'price': product.price,
                'stock': product.stock
            })
    
    return jsonify({'found': False})

@app.route('/print_invoice/<int:invoice_id>')
def print_invoice(invoice_id):
    """Thermal invoice printing view"""
    invoice = Invoice.query.get_or_404(invoice_id)
    return render_template('print_invoice.html', invoice=invoice, format_lempiras=format_lempiras)

@app.route('/export_invoices')
def export_invoices():
    """Export daily invoices to Excel"""
    export_date = request.args.get('date')
    if export_date:
        export_date = datetime.strptime(export_date, '%Y-%m-%d').date()
    else:
        export_date = datetime.now().date()
    
    # Get invoices for the specified date
    start_date = datetime.combine(export_date, datetime.min.time())
    end_date = datetime.combine(export_date, datetime.max.time())
    
    invoices = Invoice.query.filter(
        Invoice.created_at.between(start_date, end_date)
    ).order_by(Invoice.created_at.desc()).all()
    
    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Facturas {export_date.strftime('%Y-%m-%d')}"
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    # Company header
    ws.merge_cells('A1:H1')
    ws['A1'] = "Car Wash Peña Blanca - Reporte de Facturas"
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = center_alignment
    
    ws.merge_cells('A2:H2')
    ws['A2'] = f"Fecha: {export_date.strftime('%d/%m/%Y')} | Peña Blanca, Cortés | Tel: 97164446"
    ws['A2'].font = Font(size=12)
    ws['A2'].alignment = center_alignment
    
    # Column headers
    headers = ['Factura', 'Cliente', 'Teléfono', 'Total', 'Método', 'Fecha', 'Hora', 'Items']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    # Data rows
    row_num = 5
    total_sales = 0
    
    for invoice in invoices:
        # Invoice items summary
        items_summary = ', '.join([f"{item.description or (item.product.name if item.product else item.service_type.name)} ({item.quantity})" 
                                 for item in invoice.items])
        
        ws.cell(row=row_num, column=1, value=invoice.invoice_number)
        ws.cell(row=row_num, column=2, value=invoice.customer.name)
        ws.cell(row=row_num, column=3, value=invoice.customer.phone)
        ws.cell(row=row_num, column=4, value=invoice.total_amount)
        ws.cell(row=row_num, column=5, value=invoice.payment_method.title())
        ws.cell(row=row_num, column=6, value=invoice.created_at.strftime('%d/%m/%Y'))
        ws.cell(row=row_num, column=7, value=invoice.created_at.strftime('%H:%M'))
        ws.cell(row=row_num, column=8, value=items_summary)
        
        # Format currency
        ws.cell(row=row_num, column=4).number_format = 'L #,##0.00'
        
        total_sales += invoice.total_amount
        row_num += 1
    
    # Summary row
    if invoices:
        row_num += 1
        ws.merge_cells(f'A{row_num}:C{row_num}')
        summary_cell = ws.cell(row=row_num, column=1, value="TOTAL DEL DÍA:")
        summary_cell.font = Font(bold=True)
        summary_cell.alignment = Alignment(horizontal="right")
        
        total_cell = ws.cell(row=row_num, column=4, value=total_sales)
        total_cell.font = Font(bold=True)
        total_cell.number_format = 'L #,##0.00'
        total_cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    
    # Adjust column widths
    column_widths = [15, 25, 15, 12, 12, 12, 8, 40]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    
    # Save to BytesIO
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    filename = f"facturas_carwash_{export_date.strftime('%Y%m%d')}.xlsx"
    
    return send_file(
        excel_file,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/adjust_stock', methods=['POST'])
def adjust_stock():
    """Adjust product stock via AJAX"""
    try:
        data = request.get_json()
        product_id = data.get('product_id')
        action = data.get('action')
        amount = int(data.get('amount', 0))
        
        if not product_id or not action or amount <= 0:
            return jsonify({'success': False, 'error': 'Datos inválidos'})
        
        product = Product.query.get_or_404(product_id)
        
        if action == 'add':
            product.stock += amount
        elif action == 'remove':
            if product.stock >= amount:
                product.stock -= amount
            else:
                return jsonify({'success': False, 'error': 'Stock insuficiente'})
        else:
            return jsonify({'success': False, 'error': 'Acción inválida'})
        
        db.session.commit()
        return jsonify({'success': True, 'new_stock': product.stock})
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)})

@app.context_processor
def utility_processor():
    """Make utility functions available in templates"""
    return dict(format_lempiras=format_lempiras, datetime=datetime)
